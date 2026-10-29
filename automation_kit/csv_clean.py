"""csv_clean: clean a CSV or XLSX file against a declarative schema.

Reads a table, normalises every cell according to a YAML/JSON schema, rejects rows that
break the schema, removes or flags duplicates, and writes a clean file plus an error report.

Row numbers in the report are record numbers as a spreadsheet shows them: the header is
row 1 and the first data row is row 2. (A CSV record with a line break inside quotes still
counts as one row.)

Exit codes: 0 = nothing rejected, 1 = at least one row rejected, 2 = unusable input or schema.
"""

from __future__ import annotations

import argparse
import codecs
import csv
import io
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import yaml


class SchemaError(Exception):
    """The schema file is invalid."""


class DataFileError(Exception):
    """The input file cannot be used (unreadable, empty, missing columns)."""


TYPES = {"string", "integer", "number", "boolean", "date"}
CASES = {"none", "lower", "upper", "title"}
TRUE_WORDS = {"true", "t", "yes", "y", "1"}
FALSE_WORDS = {"false", "f", "no", "n", "0"}
COLUMN_KEYS = {
    "name", "aliases", "type", "required", "trim", "case", "allowed", "date_formats",
    "date_output", "default", "min", "max", "max_length", "pattern", "decimal",
    "thousands", "decimals",
}
NUMBER_RE = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)")


# --------------------------------------------------------------------------- schema


@dataclass
class Column:
    name: str
    aliases: list[str] = field(default_factory=list)
    type: str = "string"
    required: bool = False
    trim: bool = True
    case: str = "none"
    allowed: list[str] | None = None
    date_formats: list[str] = field(default_factory=lambda: ["%Y-%m-%d"])
    date_output: str = "%Y-%m-%d"
    default: str | None = None
    min: Decimal | None = None
    max: Decimal | None = None
    max_length: int | None = None
    pattern: re.Pattern[str] | None = None
    decimal: str = "."
    thousands: str | None = None
    decimals: int | None = None

    def apply_case(self, value: str) -> str:
        if self.case == "lower":
            return value.lower()
        if self.case == "upper":
            return value.upper()
        if self.case == "title":
            return value.title()
        return value


@dataclass
class Schema:
    columns: list[Column]
    dedupe_key: list[str] = field(default_factory=list)
    dedupe_keep: str = "first"
    dedupe_action: str = "drop"
    extra_columns: str = "keep"
    empty_values: set[str] = field(default_factory=lambda: {""})


def _decimal_or_none(value: Any, where: str) -> Decimal | None:
    if value is None:
        return None
    try:
        return Decimal(str(value))
    except InvalidOperation as exc:
        raise SchemaError(f"{where} must be a number") from exc


def _column_from_dict(raw: Any, index: int) -> Column:
    where = f"columns[{index}]"
    if not isinstance(raw, dict) or "name" not in raw:
        raise SchemaError(f"{where} needs a 'name'")
    unknown = set(raw) - COLUMN_KEYS
    if unknown:
        raise SchemaError(f"{where} has unknown keys: {', '.join(sorted(unknown))}")
    name = str(raw["name"])
    where = f"column '{name}'"
    col = Column(name=name)
    col.type = str(raw.get("type", "string"))
    if col.type not in TYPES:
        raise SchemaError(f"{where}: type must be one of {sorted(TYPES)}")
    col.case = str(raw.get("case", "none"))
    if col.case not in CASES:
        raise SchemaError(f"{where}: case must be one of {sorted(CASES)}")
    col.aliases = [str(a) for a in raw.get("aliases", [])]
    col.required = bool(raw.get("required", False))
    col.trim = bool(raw.get("trim", True))
    if "date_formats" in raw:
        col.date_formats = [str(f) for f in raw["date_formats"]]
        if not col.date_formats:
            raise SchemaError(f"{where}: date_formats must not be empty")
    col.date_output = str(raw.get("date_output", "%Y-%m-%d"))
    col.default = None if raw.get("default") is None else str(raw["default"])
    col.min = _decimal_or_none(raw.get("min"), f"{where}: min")
    col.max = _decimal_or_none(raw.get("max"), f"{where}: max")
    if raw.get("max_length") is not None:
        col.max_length = int(raw["max_length"])
    if raw.get("pattern") is not None:
        try:
            col.pattern = re.compile(str(raw["pattern"]))
        except re.error as exc:
            raise SchemaError(f"{where}: bad pattern: {exc}") from exc
    col.decimal = str(raw.get("decimal", "."))
    col.thousands = None if raw.get("thousands") is None else str(raw["thousands"])
    if col.decimal not in {".", ","}:
        raise SchemaError(f"{where}: decimal must be '.' or ','")
    if col.thousands is not None and col.thousands == col.decimal:
        raise SchemaError(f"{where}: thousands and decimal separators must differ")
    if raw.get("decimals") is not None:
        col.decimals = int(raw["decimals"])
        if col.decimals < 0:
            raise SchemaError(f"{where}: decimals must be >= 0")
    if raw.get("allowed") is not None:
        allowed = raw["allowed"]
        if not isinstance(allowed, list):
            raise SchemaError(f"{where}: allowed must be a list")
        col.allowed = [col.apply_case(str(v).strip() if col.trim else str(v)) for v in allowed]
    return col


def parse_schema(raw: Any) -> Schema:
    if not isinstance(raw, dict) or not isinstance(raw.get("columns"), list) or not raw["columns"]:
        raise SchemaError("schema needs a non-empty 'columns' list")
    columns = [_column_from_dict(c, i) for i, c in enumerate(raw["columns"])]
    names = [c.name for c in columns]
    if len(set(names)) != len(names):
        raise SchemaError("column names must be unique")
    schema = Schema(columns=columns)
    dedupe = raw.get("dedupe")
    if dedupe:
        key = dedupe.get("key")
        if not isinstance(key, list) or not key:
            raise SchemaError("dedupe.key must be a non-empty list of column names")
        missing = [k for k in key if k not in names]
        if missing:
            raise SchemaError(f"dedupe.key names unknown columns: {', '.join(missing)}")
        schema.dedupe_key = [str(k) for k in key]
        schema.dedupe_keep = str(dedupe.get("keep", "first"))
        schema.dedupe_action = str(dedupe.get("action", "drop"))
        if schema.dedupe_keep not in {"first", "last"}:
            raise SchemaError("dedupe.keep must be 'first' or 'last'")
        if schema.dedupe_action not in {"drop", "flag"}:
            raise SchemaError("dedupe.action must be 'drop' or 'flag'")
    schema.extra_columns = str(raw.get("extra_columns", "keep"))
    if schema.extra_columns not in {"keep", "drop", "error"}:
        raise SchemaError("extra_columns must be 'keep', 'drop' or 'error'")
    if raw.get("empty_values") is not None:
        schema.empty_values = {str(v).strip().casefold() for v in raw["empty_values"]} | {""}
    return schema


def load_schema(path: str | Path) -> Schema:
    path = Path(path)
    try:
        text = path.read_text(encoding="utf-8")
        raw = json.loads(text) if path.suffix.lower() == ".json" else yaml.safe_load(text)
    except (OSError, ValueError, yaml.YAMLError) as exc:
        raise SchemaError(f"cannot read schema {path.name}: {exc}") from exc
    return parse_schema(raw)


# ------------------------------------------------------------------------ normalising


def parse_decimal(text: str, decimal: str = ".", thousands: str | None = None) -> Decimal:
    t = text.strip()
    if thousands:
        t = t.replace(thousands, "")
    if decimal != ".":
        if "." in t:
            raise ValueError("unexpected '.'")
        t = t.replace(decimal, ".")
    if not NUMBER_RE.fullmatch(t):
        raise ValueError("not a number")
    value = Decimal(t)
    if not value.is_finite():
        raise ValueError("not finite")
    return value


def normalize_value(raw: str, col: Column, empty_values: set[str]) -> tuple[str, str | None]:
    """Return (normalised value, error reason or None)."""
    value = raw.strip() if col.trim else raw
    if value.strip().casefold() in empty_values:
        value = ""
    if value == "":
        if col.default is not None:
            value = col.default.strip() if col.trim else col.default
        elif col.required:
            return "", "required value is missing"
        else:
            return "", None

    if col.type == "string":
        value = col.apply_case(value)
    elif col.type in {"integer", "number"}:
        try:
            number = parse_decimal(value, col.decimal, col.thousands)
        except ValueError:
            return value, f"not a valid {col.type}"
        if col.type == "integer":
            if number != number.to_integral_value():
                return value, "not a whole number"
            number = number.to_integral_value()
            if col.min is not None and number < col.min or col.max is not None and number > col.max:
                return value, _range_reason(col)
            value = str(int(number))
        else:
            if col.decimals is not None:
                try:
                    number = number.quantize(Decimal(1).scaleb(-col.decimals), rounding=ROUND_HALF_UP)
                except InvalidOperation:
                    return value, "number is too large to round"
            if col.min is not None and number < col.min or col.max is not None and number > col.max:
                return value, _range_reason(col)
            value = format(number, "f")
    elif col.type == "boolean":
        word = value.casefold()
        if word in TRUE_WORDS:
            value = "true"
        elif word in FALSE_WORDS:
            value = "false"
        else:
            return value, "not a valid boolean (use yes/no, true/false, 1/0)"
    elif col.type == "date":
        parsed: date | None = None
        for fmt in col.date_formats:
            try:
                parsed = datetime.strptime(value, fmt).date()
                break
            except ValueError:
                continue
        if parsed is None:
            return value, f"not a valid date (tried {', '.join(col.date_formats)})"
        value = parsed.strftime(col.date_output)

    if col.allowed is not None and value not in col.allowed:
        return value, f"'{value}' is not one of the allowed values"
    if col.max_length is not None and len(value) > col.max_length:
        return value, f"longer than {col.max_length} characters"
    if col.pattern is not None and not col.pattern.fullmatch(value):
        return value, "does not match the required pattern"
    return value, None


def _range_reason(col: Column) -> str:
    low = "-inf" if col.min is None else format(col.min, "f")
    high = "inf" if col.max is None else format(col.max, "f")
    return f"outside the allowed range {low} to {high}"


# ------------------------------------------------------------------------- reading


@dataclass
class Table:
    headers: list[str]
    rows: list[tuple[int, list[str]]]  # (row number, cells)
    encoding: str = "n/a"


def decode_bytes(data: bytes) -> tuple[str, str]:
    """Decode text of unknown encoding. Returns (text, label of what was used)."""
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16"), "utf-16"
    try:
        text = data.decode("utf-8-sig")
        return text, "utf-8-sig" if data.startswith(codecs.BOM_UTF8) else "utf-8"
    except UnicodeDecodeError:
        pass
    # Not valid UTF-8 as a whole. Decode line by line so a file that mixes UTF-8 lines with
    # Windows-1252 lines still comes out right (UTF-8 first, then cp1252, then latin-1).
    pieces: list[str] = []
    used: set[str] = set()
    for line in data.splitlines(keepends=True):
        try:
            pieces.append(line.decode("utf-8"))
            if not line.isascii():
                used.add("utf-8")
            continue
        except UnicodeDecodeError:
            pass
        try:
            pieces.append(line.decode("cp1252"))
            used.add("cp1252")
        except UnicodeDecodeError:
            pieces.append(line.decode("latin-1"))
            used.add("latin-1")
    label = "+".join(sorted(used)) if used else "utf-8"
    return "".join(pieces), ("mixed " + label if len(used) > 1 else label)


def _read_csv(path: Path, delimiter: str, encoding: str | None) -> Table:
    data = path.read_bytes()
    if encoding:
        try:
            text, label = data.decode(encoding), encoding
        except (UnicodeDecodeError, LookupError) as exc:
            raise DataFileError(f"cannot decode {path.name} as {encoding}: {exc}") from exc
    else:
        text, label = decode_bytes(data)
    text = text.lstrip("﻿")
    if not text.strip():
        raise DataFileError(f"{path.name} is empty")
    if delimiter == "auto":
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    csv.field_size_limit(10_000_000)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    records: list[tuple[int, list[str]]] = []
    try:
        for number, cells in enumerate(reader, start=1):
            records.append((number, cells))
    except csv.Error as exc:
        raise DataFileError(f"cannot parse {path.name}: {exc}") from exc
    headers = records[0][1]
    return Table(headers=headers, rows=records[1:], encoding=label)


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d" if value.time() == datetime.min.time() else "%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_xlsx(path: Path, sheet: str | None) -> Table:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(path, data_only=True)
    except (InvalidFileException, OSError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise DataFileError(f"cannot open {path.name} as an XLSX workbook: {exc}") from exc
    if sheet is not None and sheet not in workbook.sheetnames:
        raise DataFileError(f"sheet '{sheet}' not found; available: {', '.join(workbook.sheetnames)}")
    worksheet = workbook[sheet] if sheet else workbook.active
    rows = [
        [_cell_to_text(c) for c in row]
        for row in worksheet.iter_rows(min_row=1, values_only=True)
    ]
    if not rows or not any(cell.strip() for cell in rows[0]):
        raise DataFileError(f"{path.name}: row 1 must be the header row")
    width = max(i for i, h in enumerate(rows[0]) if h.strip()) + 1
    numbered = [(n, cells[:width]) for n, cells in enumerate(rows, start=1)]
    return Table(headers=numbered[0][1], rows=numbered[1:], encoding="xlsx")


def load_table(path: str | Path, sheet: str | None = None, delimiter: str = ",", encoding: str | None = None) -> Table:
    path = Path(path)
    if not path.exists():
        raise DataFileError(f"{path} does not exist")
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path, sheet)
    if suffix == ".xls":
        raise DataFileError("legacy .xls files are not supported; save as .xlsx or .csv")
    return _read_csv(path, delimiter, encoding)


# ------------------------------------------------------------------------ cleaning


@dataclass
class Issue:
    row: int
    column: str
    value: str
    reason: str
    level: str  # "error" (row rejected), "duplicate" (row dropped), "flagged" (row kept)


@dataclass
class CleanResult:
    columns: list[str]
    rows: list[dict[str, str]]
    row_numbers: list[int]
    issues: list[Issue]
    stats: dict[str, Any]


def clean_table(table: Table, schema: Schema) -> CleanResult:
    normalized = [h.strip().casefold() for h in table.headers]
    if len(set(normalized)) != len(normalized):
        dupes = sorted({h for h in normalized if normalized.count(h) > 1})
        raise DataFileError(f"duplicate column headers: {', '.join(dupes)}")
    position = {h: i for i, h in enumerate(normalized)}

    column_index: dict[str, int | None] = {}
    missing_required: list[str] = []
    for col in schema.columns:
        found = None
        for candidate in [col.name, *col.aliases]:
            found = position.get(candidate.strip().casefold())
            if found is not None:
                break
        column_index[col.name] = found
        if found is None and col.required and col.default is None:
            missing_required.append(col.name)
    if missing_required:
        raise DataFileError(f"required columns not found in the file: {', '.join(missing_required)}")

    used = {i for i in column_index.values() if i is not None}
    extras = [i for i in range(len(table.headers)) if i not in used]
    if schema.extra_columns == "error" and extras:
        names = ", ".join(table.headers[i] for i in extras)
        raise DataFileError(f"unexpected columns in the file: {names}")
    schema_names = {c.name for c in schema.columns}
    extra_names: dict[int, str] = {}
    if schema.extra_columns == "keep":
        for i in extras:
            name = table.headers[i].strip() or f"column_{i + 1}"
            extra_names[i] = f"extra_{name}" if name in schema_names else name

    issues: list[Issue] = []
    valid: list[tuple[int, dict[str, str]]] = []
    blank = 0
    input_rows = 0
    rejected = 0
    for number, cells in table.rows:
        if not any(c.strip() for c in cells):
            blank += 1
            continue
        input_rows += 1
        if len(cells) != len(table.headers):
            issues.append(Issue(number, "", "", f"row has {len(cells)} cells, header has {len(table.headers)}", "error"))
            rejected += 1
            continue
        out: dict[str, str] = {}
        bad = False
        for col in schema.columns:
            idx = column_index[col.name]
            raw = cells[idx] if idx is not None else ""
            value, error = normalize_value(raw, col, schema.empty_values)
            if error:
                issues.append(Issue(number, col.name, raw, error, "error"))
                bad = True
            out[col.name] = value
        if bad:
            rejected += 1
            continue
        for i, name in extra_names.items():
            out[name] = cells[i]
        valid.append((number, out))

    dropped: set[int] = set()
    flagged = 0
    if schema.dedupe_key:
        groups: dict[tuple[str, ...], list[int]] = {}
        for pos, (_, out) in enumerate(valid):
            key = tuple(out[k] for k in schema.dedupe_key)
            if any(v == "" for v in key):
                continue  # a row without a key cannot be identified as a duplicate
            groups.setdefault(key, []).append(pos)
        for key, positions in groups.items():
            if len(positions) < 2:
                continue
            kept = positions[0] if schema.dedupe_keep == "first" else positions[-1]
            for pos in positions:
                if pos == kept:
                    continue
                label = ", ".join(f"{k}={v}" for k, v in zip(schema.dedupe_key, key))
                reason = f"duplicate of row {valid[kept][0]} ({label})"
                if schema.dedupe_action == "drop":
                    dropped.add(pos)
                    issues.append(Issue(valid[pos][0], ",".join(schema.dedupe_key), "", reason, "duplicate"))
                else:
                    flagged += 1
                    issues.append(Issue(valid[pos][0], ",".join(schema.dedupe_key), "", reason, "flagged"))

    kept_rows = [(n, o) for pos, (n, o) in enumerate(valid) if pos not in dropped]
    columns = [c.name for c in schema.columns] + list(extra_names.values())
    issues.sort(key=lambda i: (i.row, i.column))
    stats = {
        "input_rows": input_rows,
        "clean_rows": len(kept_rows),
        "rejected_rows": rejected,
        "duplicates_dropped": len(dropped),
        "duplicates_flagged": flagged,
        "blank_rows_skipped": blank,
        "encoding": table.encoding,
    }
    return CleanResult(
        columns=columns,
        rows=[o for _, o in kept_rows],
        row_numbers=[n for n, _ in kept_rows],
        issues=issues,
        stats=stats,
    )


# ------------------------------------------------------------------------- writing


def write_clean(path: str | Path, result: CleanResult) -> None:
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(result.columns)
        for row in result.rows:
            sheet.append([row[c] for c in result.columns])
        workbook.save(path)
        return
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(result.columns)
        for row in result.rows:
            writer.writerow([row[c] for c in result.columns])


def write_issues(path: str | Path, issues: list[Issue]) -> None:
    with Path(path).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["row", "column", "value", "reason", "level"])
        for issue in issues:
            writer.writerow([issue.row, issue.column, issue.value[:200], issue.reason, issue.level])


# ----------------------------------------------------------------------------- CLI


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="csv_clean", description="Clean a CSV/XLSX file against a schema.")
    p.add_argument("input", help="CSV or XLSX file")
    p.add_argument("--schema", required=True, help="schema file (.yaml, .yml or .json)")
    p.add_argument("--out", help="clean output (.csv or .xlsx); default <input>.clean.csv")
    p.add_argument("--errors", help="error report CSV; default <input>.errors.csv")
    p.add_argument("--sheet", help="XLSX sheet name (default: the active sheet)")
    p.add_argument("--delimiter", default=",", help="CSV delimiter, or 'auto' (default ',')")
    p.add_argument("--encoding", help="force an encoding instead of detecting it")
    p.add_argument("--strict", action="store_true", help="also exit 1 when duplicates were dropped")
    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    source = Path(args.input)
    out = Path(args.out) if args.out else source.with_name(source.stem + ".clean.csv")
    errors = Path(args.errors) if args.errors else source.with_name(source.stem + ".errors.csv")
    try:
        schema = load_schema(args.schema)
        table = load_table(source, args.sheet, args.delimiter, args.encoding)
        result = clean_table(table, schema)
    except (SchemaError, DataFileError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    write_clean(out, result)
    write_issues(errors, result.issues)
    print(json.dumps({**result.stats, "clean_file": str(out), "error_report": str(errors)}, indent=2))
    stats = result.stats
    if stats["rejected_rows"] > 0 or (args.strict and stats["duplicates_dropped"] > 0):
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
