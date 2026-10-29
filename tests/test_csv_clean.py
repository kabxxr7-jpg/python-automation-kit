import csv
import json
import subprocess
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from automation_kit import csv_clean
from automation_kit.csv_clean import (
    DataFileError,
    SchemaError,
    clean_table,
    decode_bytes,
    load_table,
    normalize_value,
    parse_schema,
)

ROOT = Path(__file__).resolve().parents[1]


def run(tmp_path, csv_text, columns, encoding="utf-8", **schema_extra):
    path = tmp_path / "in.csv"
    path.write_bytes(csv_text.encode(encoding) if isinstance(csv_text, str) else csv_text)
    schema = parse_schema({"columns": columns, **schema_extra})
    return clean_table(load_table(path), schema)


def col(**kw):
    return parse_schema({"columns": [{"name": "v", **kw}]}).columns[0]


def norm(raw, **kw):
    return normalize_value(raw, col(**kw), {""})


# ------------------------------------------------------------------ the example file


def test_example_file_end_to_end(tmp_path):
    out, errors = tmp_path / "clean.csv", tmp_path / "errors.csv"
    code = csv_clean.main([
        str(ROOT / "examples/customers.csv"), "--schema", str(ROOT / "examples/schema.yaml"),
        "--out", str(out), "--errors", str(errors),
    ])
    assert code == 1  # rows were rejected
    clean = list(csv.DictReader(out.open(encoding="utf-8")))
    assert [r["customer_id"] for r in clean] == ["101", "102", "106"]
    assert clean[0] == {
        "customer_id": "101", "full_name": "Ada Lovelace", "email": "ada@example.com",
        "signup_date": "2024-01-15", "plan": "pro", "monthly_spend": "1200.50", "active": "true",
    }
    assert clean[1]["signup_date"] == "2024-02-15" and clean[1]["active"] == "false"
    assert clean[2]["signup_date"] == "2024-05-03"  # 03/05/2024 read day-first
    report = list(csv.DictReader(errors.open(encoding="utf-8")))
    got = {(int(r["row"]), r["column"], r["level"]) for r in report}
    assert got == {
        (4, "email", "error"), (5, "signup_date", "error"), (6, "full_name", "error"),
        (7, "customer_id", "duplicate"), (9, "plan", "error"),
    }


# ------------------------------------------------------------------------ exit codes


def test_exit_code_zero_when_everything_is_clean(tmp_path):
    src = tmp_path / "ok.csv"
    src.write_text("a\n1\n2\n", encoding="utf-8")
    schema = tmp_path / "s.json"
    schema.write_text(json.dumps({"columns": [{"name": "a", "type": "integer", "required": True}]}))
    code = csv_clean.main([str(src), "--schema", str(schema), "--out", str(tmp_path / "o.csv"), "--errors", str(tmp_path / "e.csv")])
    assert code == 0
    assert (tmp_path / "e.csv").read_text().strip() == "row,column,value,reason,level"


def test_exit_code_two_for_missing_required_column(tmp_path, capsys):
    src = tmp_path / "x.csv"
    src.write_text("other\n1\n", encoding="utf-8")
    schema = tmp_path / "s.json"
    schema.write_text(json.dumps({"columns": [{"name": "email", "required": True}]}))
    assert csv_clean.main([str(src), "--schema", str(schema)]) == 2
    assert "required columns not found" in capsys.readouterr().err


@pytest.mark.parametrize("bad", [
    {"columns": []},
    {"columns": [{"name": "a", "type": "money"}]},
    {"columns": [{"name": "a"}, {"name": "a"}]},
    {"columns": [{"name": "a", "colour": "red"}]},
    {"columns": [{"name": "a", "pattern": "("}]},
    {"columns": [{"name": "a"}], "dedupe": {"key": ["b"]}},
    {"columns": [{"name": "a", "type": "number", "decimal": ",", "thousands": ","}]},
])
def test_bad_schema_is_rejected(bad):
    with pytest.raises(SchemaError):
        parse_schema(bad)


def test_exit_code_two_for_bad_schema_and_missing_file(tmp_path):
    schema = tmp_path / "s.json"
    schema.write_text("{not json")
    assert csv_clean.main([str(tmp_path / "nope.csv"), "--schema", str(schema)]) == 2
    good = tmp_path / "g.json"
    good.write_text(json.dumps({"columns": [{"name": "a"}]}))
    assert csv_clean.main([str(tmp_path / "nope.csv"), "--schema", str(good)]) == 2


def test_strict_makes_dropped_duplicates_fail(tmp_path):
    src = tmp_path / "d.csv"
    src.write_text("k\nx\nx\n", encoding="utf-8")
    schema = tmp_path / "s.json"
    schema.write_text(json.dumps({"columns": [{"name": "k"}], "dedupe": {"key": ["k"]}}))
    args = [str(src), "--schema", str(schema), "--out", str(tmp_path / "o.csv"), "--errors", str(tmp_path / "e.csv")]
    assert csv_clean.main(args) == 0
    assert csv_clean.main(args + ["--strict"]) == 1


def test_runs_as_a_module(tmp_path):
    proc = subprocess.run(
        [sys.executable, "-m", "automation_kit.csv_clean", str(ROOT / "examples/customers.csv"),
         "--schema", str(ROOT / "examples/schema.yaml"), "--out", str(tmp_path / "o.csv"),
         "--errors", str(tmp_path / "e.csv")],
        cwd=ROOT, capture_output=True, text=True,
    )
    assert proc.returncode == 1
    assert json.loads(proc.stdout)["clean_rows"] == 3


# --------------------------------------------------------------------- value rules


def test_trim_case_and_empty_values():
    assert norm("  HeLLo  ", case="lower") == ("hello", None)
    assert norm("mary o'neil", case="title") == ("Mary O'Neil", None)
    assert norm("  x  ", trim=False) == ("  x  ", None)
    assert norm("N/A", required=True) == ("N/A", None)  # not an empty marker by default
    assert normalize_value("N/A", col(required=True), {"", "n/a"}) == ("", "required value is missing")


def test_required_missing_and_default():
    assert norm("   ", required=True) == ("", "required value is missing")
    assert norm("", required=True, default="free") == ("free", None)
    assert norm("", default="0", type="integer") == ("0", None)
    assert norm("") == ("", None)


@pytest.mark.parametrize("raw,expected", [
    ("2024-02-29", "2024-02-29"), ("2023-12-31", "2023-12-31"),
])
def test_valid_dates(raw, expected):
    assert norm(raw, type="date") == (expected, None)


@pytest.mark.parametrize("raw", ["2024-02-30", "2023-02-29", "2024-13-01", "31/12/2024", "yesterday", "2024-1-1x",
                                "2024-05-03 and more", "2024-05-03T10:00:00"])
def test_bad_dates_are_rejected(raw):
    value, error = norm(raw, type="date")
    assert error and error.startswith("not a valid date")


def test_date_formats_are_tried_in_order_and_output_is_configurable():
    kw = dict(type="date", date_formats=["%d/%m/%Y", "%m/%d/%Y"], date_output="%d %b %Y")
    assert norm("03/05/2024", **kw) == ("03 May 2024", None)  # day-first wins when both fit
    assert norm("05/13/2024", **kw) == ("13 May 2024", None)  # falls through to month-first
    assert norm("13/13/2024", **kw)[1] is not None


@pytest.mark.parametrize("raw,kw,expected", [
    ("1,234", dict(type="integer", thousands=","), "1234"),
    ("12.0", dict(type="integer"), "12"),
    ("-7", dict(type="integer"), "-7"),
    ("1.234,50", dict(type="number", decimal=",", thousands="."), "1234.50"),
    ("1,234.5", dict(type="number", thousands=","), "1234.5"),
    ("2.675", dict(type="number", decimals=2), "2.68"),  # half up on the decimal, not on a float
    ("2.665", dict(type="number", decimals=2), "2.67"),  # half-even would give 2.66
    ("0.125", dict(type="number", decimals=2), "0.13"),  # half-even would give 0.12
    ("10.5", dict(type="number", decimals=2), "10.50"),
    (".5", dict(type="number"), "0.5"),
])
def test_numbers(raw, kw, expected):
    assert norm(raw, **kw) == (expected, None)


@pytest.mark.parametrize("raw,kw", [
    ("12.5", dict(type="integer")),
    ("abc", dict(type="number")),
    ("1e5", dict(type="number")),
    ("NaN", dict(type="number")),
    ("Infinity", dict(type="number")),
    ("1_000", dict(type="integer")),
    ("1.5", dict(type="number", decimal=",")),  # a '.' where ',' is the decimal mark is ambiguous
    ("-1", dict(type="integer", min=0)),
    ("11", dict(type="number", max=10)),
])
def test_bad_numbers(raw, kw):
    assert norm(raw, **kw)[1] is not None


@pytest.mark.parametrize("raw,expected", [
    ("Yes", "true"), ("y", "true"), ("TRUE", "true"), ("1", "true"),
    ("no", "false"), ("N", "false"), ("0", "false"),
])
def test_booleans(raw, expected):
    assert norm(raw, type="boolean") == (expected, None)
    assert norm("maybe", type="boolean")[1] is not None


def test_allowed_pattern_and_length():
    assert norm("PRO", case="lower", allowed=["Free", "PRO"]) == ("pro", None)  # allowed list gets the same case rule
    assert norm("gold", case="lower", allowed=["free", "pro"])[1]
    assert norm("a@b.co", pattern=r"[^@\s]+@[^@\s]+\.\w+") == ("a@b.co", None)
    assert norm("a@b", pattern=r"[^@\s]+@[^@\s]+\.\w+")[1]
    assert norm("xxxxxx", max_length=5)[1] == "longer than 5 characters"


# ----------------------------------------------------------------- table handling


def test_aliases_and_header_matching_ignore_case_and_spaces(tmp_path):
    result = run(tmp_path, " E-Mail ,NAME\na@b.co,x\n", [
        {"name": "email", "aliases": ["e-mail"]}, {"name": "name"},
    ])
    assert result.rows == [{"email": "a@b.co", "name": "x"}]


def test_extra_columns_modes(tmp_path):
    cols = [{"name": "a"}]
    kept = run(tmp_path, "a,note\n1,hi\n", cols, extra_columns="keep")
    assert kept.columns == ["a", "note"] and kept.rows[0]["note"] == "hi"
    dropped = run(tmp_path, "a,note\n1,hi\n", cols, extra_columns="drop")
    assert dropped.columns == ["a"]
    with pytest.raises(DataFileError, match="unexpected columns"):
        run(tmp_path, "a,note\n1,hi\n", cols, extra_columns="error")


def test_duplicate_headers_are_fatal(tmp_path):
    with pytest.raises(DataFileError, match="duplicate column headers"):
        run(tmp_path, "a,A\n1,2\n", [{"name": "a"}])


def test_optional_missing_column_gets_default(tmp_path):
    result = run(tmp_path, "a\n1\n", [{"name": "a"}, {"name": "b", "default": "n"}])
    assert result.rows == [{"a": "1", "b": "n"}]


def test_row_numbers_header_is_row_one_and_blank_lines_are_skipped(tmp_path):
    result = run(tmp_path, "a\n1\n\nx\n3\n,\n", [{"name": "a", "type": "integer"}])
    assert result.stats["blank_rows_skipped"] == 2
    assert [(i.row, i.column) for i in result.issues] == [(4, "a")]  # 'x' is on record 4
    assert result.row_numbers == [2, 5]


def test_quoted_newline_counts_as_one_row(tmp_path):
    result = run(tmp_path, 'a,b\n"line one\nline two",1\nbad,x\n', [{"name": "a"}, {"name": "b", "type": "integer"}])
    assert result.rows[0]["a"] == "line one\nline two"
    assert result.issues[0].row == 3


def test_wrong_cell_count_rejects_the_row(tmp_path):
    result = run(tmp_path, "a,b\n1,2\n3\n4,5,6\n", [{"name": "a"}, {"name": "b"}])
    assert len(result.rows) == 1
    assert [(i.row, i.reason) for i in result.issues] == [
        (3, "row has 1 cells, header has 2"), (4, "row has 3 cells, header has 2"),
    ]


def test_auto_delimiter(tmp_path):
    path = tmp_path / "semi.csv"
    path.write_text("a;b\n1;2\n", encoding="utf-8")
    table = load_table(path, delimiter="auto")
    assert table.headers == ["a", "b"] and table.rows == [(2, ["1", "2"])]


def test_empty_file_is_fatal(tmp_path):
    path = tmp_path / "e.csv"
    path.write_text("  \n", encoding="utf-8")
    with pytest.raises(DataFileError, match="empty"):
        load_table(path)


# --------------------------------------------------------------------- duplicates


DUP_CSV = "id,name\n1,first\n2,solo\n1,second\n1,third\n,blank\n,blank2\n"
DUP_COLS = [{"name": "id"}, {"name": "name"}]


def test_dedupe_keep_first_drops_later_rows(tmp_path):
    r = run(tmp_path, DUP_CSV, DUP_COLS, dedupe={"key": ["id"], "keep": "first", "action": "drop"})
    assert [x["name"] for x in r.rows] == ["first", "solo", "blank", "blank2"]  # empty keys are never duplicates
    dups = [i for i in r.issues if i.level == "duplicate"]
    assert [(i.row, i.reason) for i in dups] == [(4, "duplicate of row 2 (id=1)"), (5, "duplicate of row 2 (id=1)")]
    assert r.stats["duplicates_dropped"] == 2


def test_dedupe_keep_last(tmp_path):
    r = run(tmp_path, DUP_CSV, DUP_COLS, dedupe={"key": ["id"], "keep": "last", "action": "drop"})
    assert [x["name"] for x in r.rows] == ["solo", "third", "blank", "blank2"]
    assert [i.reason for i in r.issues if i.level == "duplicate"] == ["duplicate of row 5 (id=1)"] * 2


def test_dedupe_flag_keeps_every_row(tmp_path):
    r = run(tmp_path, DUP_CSV, DUP_COLS, dedupe={"key": ["id"], "action": "flag"})
    assert len(r.rows) == 6
    assert r.stats["duplicates_flagged"] == 2 and r.stats["duplicates_dropped"] == 0
    assert {i.level for i in r.issues} == {"flagged"}


def test_rejected_rows_do_not_take_part_in_dedupe(tmp_path):
    csv_text = "id,qty\n1,oops\n1,5\n"
    r = run(tmp_path, csv_text, [{"name": "id"}, {"name": "qty", "type": "integer"}], dedupe={"key": ["id"]})
    assert [x["qty"] for x in r.rows] == ["5"]  # the bad row 2 was rejected, so row 3 is not a duplicate
    assert r.stats["duplicates_dropped"] == 0


def test_dedupe_compares_normalised_values(tmp_path):
    r = run(tmp_path, "email\nA@X.COM\n a@x.com \n", [{"name": "email", "case": "lower"}], dedupe={"key": ["email"]})
    assert r.rows == [{"email": "a@x.com"}] and r.stats["duplicates_dropped"] == 1


# ---------------------------------------------------------------------- encodings


def test_utf8_bom(tmp_path):
    r = run(tmp_path, b"\xef\xbb\xbfname\nJos\xc3\xa9\n", [{"name": "name"}])
    assert r.rows == [{"name": "José"}] and r.stats["encoding"] == "utf-8-sig"


def test_windows_1252(tmp_path):
    r = run(tmp_path, "name\nJosé\nZoë\n€5\n".encode("cp1252"), [{"name": "name"}])
    assert [x["name"] for x in r.rows] == ["José", "Zoë", "€5"] and r.stats["encoding"] == "cp1252"


def test_mixed_utf8_and_cp1252_lines(tmp_path):
    data = "name\n".encode() + "José\n".encode("utf-8") + "Zoë\n".encode("cp1252") + b"plain\n"
    r = run(tmp_path, data, [{"name": "name"}])
    assert [x["name"] for x in r.rows] == ["José", "Zoë", "plain"]
    assert r.stats["encoding"] == "mixed cp1252+utf-8"


def test_undefined_cp1252_byte_falls_back_to_latin1():
    text, label = decode_bytes(b"a\n\x81\n")
    assert text == "a\n\x81\n" and label == "latin-1"


def test_utf16_with_bom(tmp_path):
    r = run(tmp_path, "name\nJosé\n".encode("utf-16"), [{"name": "name"}])
    assert r.rows == [{"name": "José"}] and r.stats["encoding"] == "utf-16"


def test_forced_encoding(tmp_path):
    path = tmp_path / "f.csv"
    path.write_bytes("name\nJosé\n".encode("cp1252"))
    assert load_table(path, encoding="cp1252").rows == [(2, ["José"])]
    with pytest.raises(DataFileError):
        load_table(path, encoding="ascii")


def test_output_is_utf8_even_when_input_was_not(tmp_path):
    src = tmp_path / "in.csv"
    src.write_bytes("name\nJosé\n".encode("cp1252"))
    schema = tmp_path / "s.json"
    schema.write_text(json.dumps({"columns": [{"name": "name"}]}))
    out = tmp_path / "o.csv"
    csv_clean.main([str(src), "--schema", str(schema), "--out", str(out), "--errors", str(tmp_path / "e.csv")])
    assert out.read_bytes() == "name\nJosé\n".encode("utf-8")


# --------------------------------------------------------------------------- XLSX


def make_xlsx(path, rows, sheet_title="Data"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_xlsx_row_numbers_match_excel_and_dates_and_numbers(tmp_path):
    path = tmp_path / "in.xlsx"
    make_xlsx(path, [
        ["id", "when", "amount", "ok"],
        [1, datetime(2024, 3, 9), 12.0, True],
        [None, None, None, None],          # Excel row 3 is blank
        [2, "not a date", 7.5, False],     # Excel row 4
        [3, datetime(2024, 3, 10, 14, 30), "x", True],   # Excel row 5
    ])
    cols = [
        {"name": "id", "type": "integer"},
        {"name": "when", "type": "date", "date_formats": ["%Y-%m-%d", "%Y-%m-%d %H:%M:%S"]},
        {"name": "amount", "type": "number"},
        {"name": "ok", "type": "boolean"},
    ]
    result = clean_table(load_table(path), parse_schema({"columns": cols}))
    assert result.rows == [{"id": "1", "when": "2024-03-09", "amount": "12", "ok": "true"}]
    assert sorted((i.row, i.column) for i in result.issues) == [(4, "when"), (5, "amount")]
    assert result.stats["blank_rows_skipped"] == 1 and result.stats["encoding"] == "xlsx"


def test_xlsx_sheet_selection_and_output(tmp_path):
    path = tmp_path / "in.xlsx"
    wb = Workbook()
    wb.active.title = "First"
    wb.active.append(["a"])
    wb.active.append(["ignored"])
    second = wb.create_sheet("Second")
    second.append(["a"])
    second.append(["picked"])
    wb.save(path)
    schema = tmp_path / "s.json"
    schema.write_text(json.dumps({"columns": [{"name": "a"}]}))
    out = tmp_path / "o.xlsx"
    code = csv_clean.main([str(path), "--schema", str(schema), "--sheet", "Second", "--out", str(out), "--errors", str(tmp_path / "e.csv")])
    assert code == 0
    assert [[c.value for c in row] for row in load_workbook(out).active.iter_rows()] == [["a"], ["picked"]]
    assert csv_clean.main([str(path), "--schema", str(schema), "--sheet", "Nope"]) == 2


def test_legacy_xls_and_corrupt_xlsx_are_reported(tmp_path):
    xls = tmp_path / "old.xls"
    xls.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DataFileError, match="legacy"):
        load_table(xls)
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"this is not a zip")
    with pytest.raises(DataFileError):
        load_table(bad)
